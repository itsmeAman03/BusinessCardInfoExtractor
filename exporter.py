"""CSV / XLSX / JSONL export of extracted business cards."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import FIELDS

HEADERS = ["Image File"] + FIELDS

_COLUMN_WIDTHS = {
    "Image File": 28,
    "First Name": 18,
    "Last Name": 18,
    "Position / Job Title": 26,
    "Company": 24,
    "Location": 44,
    "Phone Number": 32,
    "Email Address": 34,
}


def _rows(results):
    for card in results:
        data = card.get("data") or {}
        row = {"Image File": card.get("filename", "")}
        for field in FIELDS:
            row[field] = data.get(field, "Null")
        yield row


def save_csv(path, results) -> None:
    """UTF-8 with BOM so Excel opens it correctly (no mojibake in Windows)."""
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(_rows(results))


def save_jsonl(path, results) -> None:
    path = Path(path)
    with path.open("w", encoding="utf-8") as fh:
        for row in _rows(results):
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def save_xlsx(path, results) -> None:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Business Cards"

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    wrap_columns = {"Location", "Company", "Position / Job Title", "Image File"}
    for r, row in enumerate(_rows(results), start=2):
        for c, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=r, column=c, value=row[header])
            cell.alignment = Alignment(vertical="top", wrap_text=header in wrap_columns)

    for c, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = _COLUMN_WIDTHS.get(
            header, 20
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)
